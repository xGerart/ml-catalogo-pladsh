"""Excel de revisión del etiquetado asistido y construcción del dataset publicado."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from catalogo.taxonomy import NO_SE, OTROS

COLUMNS = ["id", "nombre", "propuesta", "final"]
MIN_STRATIFIED = 13  # tras el 80/20 quedan ≥ 10 por clase en train: lo que pide la CV de 10 folds


def write_review_workbook(df: pd.DataFrame, categories: list[str], path: Path) -> None:
    """Excel ordenado por propuesta, con `final` prellenado y desplegable cerrado."""
    rows = df.sort_values(["propuesta", "nombre"])
    wb = Workbook()
    ws = wb.active
    ws.title = "revision"
    ws.append(COLUMNS)
    for r in rows.itertuples(index=False):
        ws.append([r.id, r.nombre, r.propuesta, r.propuesta])
    opciones = ",".join([*categories, NO_SE])
    dv = DataValidation(type="list", formula1=f'"{opciones}"', allow_blank=False)
    dv.error = "Elige una categoría de la lista"
    ws.add_data_validation(dv)
    dv.add(f"D2:D{len(rows) + 1}")
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 18
    ws.freeze_panes = "A2"
    wb.save(path)


def read_review_workbook(path: Path) -> pd.DataFrame:
    ws = load_workbook(path, read_only=True).active
    values = list(ws.values)
    return pd.DataFrame(values[1:], columns=list(values[0]))[COLUMNS]


def build_labeled_dataset(
    review: pd.DataFrame, categories: list[str], min_count: int = 30
) -> tuple[pd.DataFrame, dict]:
    """Valida la revisión, excluye «No sé», fusiona clases chicas en «Otros»."""
    vacias = review.loc[review["final"].isna() | (review["final"].astype(str).str.strip() == ""), "id"]
    if len(vacias):
        raise ValueError(f"Revisión incompleta: celdas `final` vacías en ids {vacias.tolist()}")
    validas = set(categories) | {NO_SE}
    fuera = sorted(set(review["final"]) - validas)
    if fuera:
        raise ValueError(f"Etiquetas fuera de la taxonomía: {fuera}")

    no_se = review["final"] == NO_SE
    df = review[~no_se].copy()
    corregidos = int((df["final"] != df["propuesta"]).sum())

    counts = df["final"].value_counts()
    chicas = sorted(c for c, n in counts.items() if n < min_count and c != OTROS)
    df["etiqueta"] = df["final"].where(~df["final"].isin(chicas), OTROS)

    final_counts = df["etiqueta"].value_counts()
    insuficientes = final_counts[final_counts < MIN_STRATIFIED]
    if len(insuficientes):
        raise ValueError(
            f"Clases con menos de {MIN_STRATIFIED} ejemplos tras fusionar: {insuficientes.to_dict()}"
        )

    resumen = {
        "revisados": int(len(review)),
        "excluidos_no_se": int(no_se.sum()),
        "etiquetados": int(len(df)),
        "corregidos": corregidos,
        "pct_corregidos": round(100 * corregidos / len(df), 2),
        "fusionadas_en_otros": chicas,
        "distribucion": {k: int(v) for k, v in final_counts.items()},
    }
    out = df[["id", "nombre", "etiqueta"]].sort_values("id").reset_index(drop=True)
    return out, resumen


def load_labeled(path: Path) -> pd.DataFrame:
    return pd.read_csv(path, dtype={"nombre": str, "etiqueta": str})
